"""Build the milestone-2 Word report from report.md.

Usage (from anywhere):  python -I build_report.py
Output: ../NudgeWell_Group6_Literature_Data_Technology_Review.docx

report.md conventions
- Front matter between the first two '---' lines (title, subtitle, tagline, course, group, members, date).
- '# ', '## ', '### ' headings; '- ' bullets; '1. ' numbered items; '> ' summary box.
- 'Table N. caption' + optional '<!-- widths: a,b,c -->' + a pipe table; a first cell starting with '§' is a theme row.
- '![Figure N. caption](relative/path.png)' figures; '[[TOC]]' and '[[PAGEBREAK]]' markers.
- Paragraphs under a heading that starts with 'References' get APA hanging indents.
- Inline **bold**, *italic*, `code`, and bare https:// links.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE / "report.md"
OUT = HERE.parent / "NudgeWell_Group6_Literature_Data_Technology_Review.docx"

BODY_FONT, HEAD_FONT, CODE_FONT = "Aptos", "Aptos Display", "Consolas"
NAVY, INK, GRAY, LINK = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x22, 0x22, 0x22), RGBColor(0x55, 0x55, 0x55), "1F5FAD"
HEAD_FILL, THEME_FILL, BOX_FILL, RULE = "E7EEF8", "F3F6FB", "F7F9FC", "C9D3E0"
TEXT_WIDTH_CM = 17.0  # A4 with 2 cm margins

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|`[^`]+`|https?://[^\s)]+)")


# ---------- low-level XML helpers ----------

def set_rfonts(rpr, name):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def strip_theme_color(rpr):
    color = rpr.find(qn("w:color"))
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            color.attrib.pop(qn(attr), None)


# Schema order matters to Word: insert before the elements that must follow.
PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                  "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
                  "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
                  "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
                  "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
TBLPR_ORDER = ("w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


def shade(element_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    successors = PPR_AFTER_PBDR[1:] if element_pr.tag == qn("w:pPr") else TCPR_AFTER_SHD
    element_pr.insert_element_before(shd, *successors)


def tblpr_insert(tblpr, el):
    tag = el.tag.split("}")[1]
    later = TBLPR_ORDER[TBLPR_ORDER.index(f"w:{tag}") + 1:]
    tblpr.insert_element_before(el, *later)


def border(parent, side, size=4, color=RULE, space=0):
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), str(color))  # RGBColor prints as hex
    parent.append(el)


def add_field(paragraph, instr, placeholder=""):
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
        elif kind == "text":
            el = OxmlElement("w:t")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def add_hyperlink(paragraph, url, size):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK)
    rpr.append(color)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


# ---------- text ----------

def add_run(paragraph, text, size=None, bold=False, italic=False, color=None, font=None):
    run = paragraph.add_run(text)
    run.bold, run.italic = bold or None, italic or None
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if font:
        set_rfonts(run._r.get_or_add_rPr(), font)
    return run


def add_inline(paragraph, text, size=None, bold=False, italic=False, color=None):
    pos = 0
    for m in TOKEN.finditer(text):
        if m.start() > pos:
            add_run(paragraph, text[pos:m.start()], size, bold, italic, color)
        tok = m.group(0)
        if tok.startswith("**"):
            add_inline(paragraph, tok[2:-2], size, True, italic, color)
        elif tok.startswith("`"):
            add_run(paragraph, tok[1:-1], (size or 10.5) * 0.9, bold, italic, color, CODE_FONT)
        elif tok.startswith("http"):
            add_hyperlink(paragraph, tok, size)
        else:
            add_inline(paragraph, tok[1:-1], size, bold, True, color)
        pos = m.end()
    if pos < len(text):
        add_run(paragraph, text[pos:], size, bold, italic, color)


def caption(doc, text, before=False):
    label, _, rest = text.partition(". ")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4 if before else 2)
    p.paragraph_format.space_after = Pt(4 if before else 12)
    p.paragraph_format.keep_with_next = before
    add_run(p, label + ". ", 9, bold=True, color=NAVY)
    add_inline(p, rest, 9, color=GRAY)


# ---------- document setup ----------

def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.9)
    sec.footer_distance = Cm(0.9)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    set_rfonts(normal.element.get_or_add_rPr(), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (("Heading 1", 20, 0, 12), ("Heading 2", 14, 14, 6), ("Heading 3", 11.5, 10, 4)):
        s = doc.styles[name]
        rpr = s.element.get_or_add_rPr()
        set_rfonts(rpr, HEAD_FONT)
        strip_theme_color(rpr)
        s.font.size, s.font.bold, s.font.italic = Pt(size), True, False
        s.font.color.rgb = NAVY
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
        s.paragraph_format.keep_with_next = True

    bullet = doc.styles["List Bullet"]
    bullet.paragraph_format.space_after = Pt(3)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "NudgeWell · Group 6 · Literature, Data & Technology Review · page ", 8.5, color=GRAY)
    add_field(footer, "PAGE", "1")
    for r in footer.runs:
        r.font.size, r.font.color.rgb = Pt(8.5), GRAY

    props = doc.core_properties
    props.title = "NudgeWell: Literature, Data & Technology Review"
    props.author = "Group 6"
    props.subject = "SIC AI Cohort 18 Digital Addiction Capstone, milestone 2"


def title_page(doc, meta):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(110)
    add_run(p, meta["title"], 36, bold=True, color=NAVY, font=HEAD_FONT)
    p = doc.add_paragraph()
    add_run(p, meta["subtitle"], 18, color=NAVY, font=HEAD_FONT)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    add_run(p, meta["tagline"], 12, italic=True, color=GRAY)
    p = doc.add_paragraph()
    border(_pbdr(p), "bottom", 8, NAVY)
    for label, key in (("", "course"), ("", "group"), ("Members:", "members"), ("Submitted:", "date")):
        q = doc.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        value = meta[key]
        if label:
            # Label column: values align at a tab stop; one member per line.
            fmt = q.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Cm(2.1), Cm(-2.1)
            fmt.tab_stops.add_tab_stop(Cm(2.1))
            add_run(q, label + "\t", 10.5, bold=True, color=NAVY)
            if key == "members":
                value = "\n".join(name.strip() for name in value.split("·"))
        add_run(q, value, 10.5, bold=(key == "group"))
    doc.add_paragraph().paragraph_format.space_after = Pt(10)


def _pbdr(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(pbdr, *PPR_AFTER_PBDR)
    return pbdr


def summary_box(doc, paragraphs):
    for i, text in enumerate(paragraphs):
        p = doc.add_paragraph()
        ppr = p._p.get_or_add_pPr()
        pbdr = _pbdr(p)
        border(pbdr, "left", 18, NAVY, 8)
        shade(ppr, BOX_FILL)
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.space_after = Pt(0 if i < len(paragraphs) - 1 else 10)
        p.paragraph_format.space_before = Pt(0)
        add_inline(p, text, 10)


def toc(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    add_run(p, "Contents", 16, bold=True, color=NAVY, font=HEAD_FONT)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")


def table(doc, rows, widths):
    ncol = len(rows[0])
    widths = widths or [1] * ncol
    cms = [w / sum(widths) * TEXT_WIDTH_CM for w in widths]
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "bottom", "insideH"):
        border(borders, side, 4, RULE)
    tblpr_insert(tblpr, borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr_insert(tblpr, layout)
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", 45), ("left", 70), ("bottom", 45), ("right", 70)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr_insert(tblpr, mar)
    for gc, cm in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), cms):  # fixed layout reads the grid
        gc.set(qn("w:w"), str(round(cm / 2.54 * 1440)))

    for i, row in enumerate(rows):
        cells = t.add_row().cells
        trpr = t.rows[-1]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        if i == 0:
            hdr = OxmlElement("w:tblHeader")
            trpr.append(hdr)
        for j, c in enumerate(cells):
            c.width = Cm(cms[j])
        theme = row[0].startswith("§")
        # Keep the header and theme rows with the row below; keep short tables together as one block.
        keep = i == 0 or theme or (len(rows) <= 9 and i < len(rows) - 1)
        if theme:
            merged = cells[0].merge(cells[-1])
            shade(merged._tc.get_or_add_tcPr(), THEME_FILL)
            p = merged.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            add_run(p, row[0].lstrip("§ ").strip(), 8.5, bold=True, color=NAVY)
            continue
        for j, text in enumerate(row):
            cell = cells[j]
            if i == 0:
                shade(cell._tc.get_or_add_tcPr(), HEAD_FILL)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = keep
            add_inline(p, text, 8.5, bold=(i == 0), color=NAVY if i == 0 else None)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def figure(doc, cap, path):
    doc.add_picture(str((HERE / path).resolve()), width=Cm(16.2))
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.paragraph_format.space_before = Pt(4)
    pic.paragraph_format.space_after = Pt(0)
    caption(doc, cap)


def numbered(doc, n, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.75)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))
    add_run(p, f"{n}.\t")
    add_inline(p, text)


def reference(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    add_inline(p, text, 9.5)


# ---------- markdown driver ----------

def parse_front(lines):
    meta, i = {}, 0
    if lines and lines[0].strip() == "---":
        i = 1
        while lines[i].strip() != "---":
            k, _, v = lines[i].partition(":")
            meta[k.strip()] = v.strip()
            i += 1
        i += 1
    return meta, lines[i:]


def build():
    meta, lines = parse_front(SRC.read_text(encoding="utf-8").splitlines())
    doc = Document()
    setup(doc)
    title_page(doc, meta)

    in_refs, i = False, 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s == "[[TOC]]":
            toc(doc)
            i += 1
            continue
        if s == "[[PAGEBREAK]]":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            in_refs = False
            i += 1
            continue
        m = re.match(r"^(#{1,3}) (.+)$", s)
        if m:
            level, text = len(m.group(1)), m.group(2)
            doc.add_heading(text, level=level)
            in_refs = text.startswith("References")
            i += 1
            continue
        if s.startswith(">"):
            paras, buf = [], []
            while i < len(lines) and lines[i].strip().startswith(">"):
                t = lines[i].strip()[1:].strip()
                if t:
                    buf.append(t)
                elif buf:
                    paras.append(" ".join(buf))
                    buf = []
                i += 1
            if buf:
                paras.append(" ".join(buf))
            summary_box(doc, paras)
            continue
        m = re.match(r"^!\[(.+)\]\((.+)\)$", s)
        if m:
            figure(doc, m.group(1), m.group(2))
            i += 1
            continue
        if re.match(r"^Table \d+\. ", s):
            caption(doc, s, before=True)
            i += 1
            widths = None
            if i < len(lines) and lines[i].strip().startswith("<!-- widths:"):
                widths = [float(x) for x in re.findall(r"[\d.]+", lines[i])]
                i += 1
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            table(doc, rows, widths)
            continue
        if s.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, lines[i].strip()[2:])
                i += 1
            continue
        if re.match(r"^\d+\. ", s):
            while i < len(lines) and re.match(r"^\d+\. ", lines[i].strip()):
                n, text = lines[i].strip().split(". ", 1)
                numbered(doc, n, text)
                i += 1
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3} |- |\d+\. |>|!\[|\|)", lines[i].strip()) \
                and not lines[i].strip().startswith("[[") and not re.match(r"^Table \d+\. ", lines[i].strip()):
            buf.append(lines[i].strip())
            i += 1
        text = " ".join(buf)
        if in_refs:
            reference(doc, text)
        else:
            add_inline(doc.add_paragraph(), text)

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
